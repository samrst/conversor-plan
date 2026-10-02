import openpyxl
import xlrd
import os
import re
import unicodedata
from collections import defaultdict

def normalize_string(s):
    if not s: return ''
    nfd = unicodedata.normalize('NFD', str(s))
    no_accents = ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]', '', no_accents.lower())

def normalize_cod_turma(cod):
    if not cod: return ''
    s = str(cod).strip()
    if not s: return ''
    s = re.sub(r'^(?:saep|[tbg])[\s_-]*', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'[\s_-]*(?:emb|emp)$', '', s, flags=re.IGNORECASE).strip()
    return s

def extract_turma_code(turma_desc):
    if not turma_desc: return ''
    s = str(turma_desc).strip()
    if not s: return ''
    raw_code = s
    for sep in ['-', '–', '—']:
        if sep in raw_code:
            raw_code = raw_code.split(sep)[0].strip()
            break
    return normalize_cod_turma(raw_code)

def get_turma_display(turma_desc):
    return extract_turma_code(turma_desc)

def is_senai_ited(escola):
    if not escola: return False
    norm = normalize_string(escola)
    return 'senaiited' in norm or norm == 'ited'

def get_unidade_operacional(escola, turma_desc):
    raw_escola = (escola or '').strip()
    if is_senai_ited(raw_escola):
        desc = (turma_desc or '').strip()
        for sep in ['-', '–', '—']:
            if sep in desc:
                parts = desc.split(sep)
                extracted = sep.join(parts[1:]).strip()
                if extracted:
                    return extracted
        return desc or raw_escola
    return raw_escola

def get_modalidade(escola):
    return 'Semipresencial' if is_senai_ited(escola) else 'Presencial'

def normalize_unidade(unidade):
    if not unidade: return ''
    s = normalize_string(unidade)
    s = re.sub(r'^(?:senai|cfp|escola|unidade|centro|cetind)+', '', s)
    canonical = {
        'feiradesantana': 'feira',
        'feiradesantanabest': 'feira',
        'feiraserrinha': 'feira',
        'vitoriadaconquista': 'conquista',
        'luiseduardomagalhaes': 'lem',
        'luiseduardo': 'lem',
        'salvadorcimatec': 'cimatec',
        'salvadordendezeiros': 'dendezeiros'
    }
    for k, v in canonical.items():
        if s.startswith(k) or s == k:
            return v
    return s

def get_canonical_curso(curso):
    if not curso: return ''
    trimmed = str(curso).strip()
    norm = normalize_string(trimmed)
    if norm in ['tecnicoemautomacao', 'automacao', 'tecnicoemautomacaoindustrial', 'automacaoindustrial']:
        return 'Técnico em Automação Industrial'
    return trimmed

def normalize_curso(curso):
    if not curso: return ''
    canonical = get_canonical_curso(curso)
    s = normalize_string(canonical)
    s = re.sub(r'^(?:tecnicoem|tecnicode|cursobtecnico|curso)+', '', s)
    return s

def make_composite_key(unidade, curso, turma_code):
    u = normalize_unidade(unidade)
    c = normalize_curso(curso)
    t = normalize_cod_turma(turma_code)
    return f"{u}___{c}___{t}"

# 1. Audit File Reading
print("=================================================================")
print("=== 1. AUDITING REAL SPREADSHEETS ===")
print("=================================================================")

# Aptos
path_aptos = r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx'
wb_aptos = openpyxl.load_workbook(path_aptos, data_only=True)
assert 'Estudante' in wb_aptos.sheetnames
ws_aptos = wb_aptos['Estudante']
headers_aptos = [ws_aptos.cell(1, c).value for c in range(1, ws_aptos.max_column + 1)]
rows_aptos = []
for r in range(2, ws_aptos.max_row + 1):
    vals = [ws_aptos.cell(r, c).value for c in range(1, ws_aptos.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_aptos.append(dict(zip(headers_aptos, vals)))

print(f"Aptos File: '{os.path.basename(path_aptos)}'")
print(f"  - Sheet: 'Estudante' (from {wb_aptos.sheetnames})")
print(f"  - Total Rows: {len(rows_aptos)}")
print(f"  - Key Headers Present: ESCOLA={'ESCOLA' in headers_aptos}, CURSO={'CURSO' in headers_aptos}, TURMADESCRICAO={'TURMADESCRICAO' in headers_aptos}")

# CHP
path_chp = r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls'
wb_chp = xlrd.open_workbook(path_chp)
assert 'Dados' in wb_chp.sheet_names()
ws_chp = wb_chp.sheet_by_name('Dados')
headers_chp = [ws_chp.cell_value(1, c) for c in range(ws_chp.ncols)]
rows_chp = []
for r in range(2, ws_chp.nrows):
    vals = [ws_chp.cell_value(r, c) for c in range(ws_chp.ncols)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_chp.append(dict(zip(headers_chp, vals)))

print(f"\nCHP File: '{os.path.basename(path_chp)}'")
print(f"  - Sheet: 'Dados' (Header at Row 2, from {wb_chp.sheet_names()})")
print(f"  - Total Rows: {len(rows_chp)}")
print(f"  - Key Headers Present: CODTURMA={'CODTURMA' in headers_chp}, UNIDADE={'UNIDADE' in headers_chp}, CURSO={'CURSO' in headers_chp}")

# IFTP
path_iftp = r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx'
wb_iftp = openpyxl.load_workbook(path_iftp, data_only=True)
assert 'Dados' in wb_iftp.sheetnames
ws_iftp = wb_iftp['Dados']
headers_iftp = [ws_iftp.cell(1, c).value for c in range(1, ws_iftp.max_column + 1)]
rows_iftp = []
for r in range(2, ws_iftp.max_row + 1):
    vals = [ws_iftp.cell(r, c).value for c in range(1, ws_iftp.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_iftp.append(dict(zip(headers_iftp, vals)))

print(f"\nIFTP File: '{os.path.basename(path_iftp)}'")
print(f"  - Sheet: 'Dados' (from {wb_iftp.sheetnames})")
print(f"  - Total Rows: {len(rows_iftp)}")
print(f"  - Key Headers Present: CODTURMA={'CODTURMA' in headers_iftp}, UNIDADE={'UNIDADE' in headers_iftp}, CURSO={'CURSO' in headers_iftp}")

# Participação
path_part = r'C:\Users\sam.sampaio\Downloads\RelatorioParticipacaoTurmasProvas_2026-06-22_14-42-12.xlsx'
wb_part = openpyxl.load_workbook(path_part, data_only=True)
ws_part = wb_part.active
headers_part = [ws_part.cell(1, c).value for c in range(1, ws_part.max_column + 1)]
rows_part = []
for r in range(2, ws_part.max_row + 1):
    vals = [ws_part.cell(r, c).value for c in range(1, ws_part.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_part.append(dict(zip(headers_part, vals)))

print(f"\nParticipação File: '{os.path.basename(path_part)}'")
print(f"  - Active Sheet: '{ws_part.title}'")
print(f"  - Total Rows: {len(rows_part)}")
print(f"  - Available Headers: {headers_part}")

# 2. Consolidation & Business Rules Audit
print("\n=================================================================")
print("=== 2. AUDITING BUSINESS RULES & CONSOLIDATION ===")
print("=================================================================")

# Indexes
index_chp = defaultdict(int)
for r in rows_chp:
    c = r.get('CODTURMA')
    u = r.get('UNIDADE') or r.get('ESCOLA') or ''
    cur = r.get('CURSO') or ''
    if c:
        can_c = get_canonical_curso(cur)
        comp = make_composite_key(u, can_c, c)
        index_chp[comp] += 1

index_iftp = defaultdict(int)
for r in rows_iftp:
    c = r.get('CODTURMA')
    u = r.get('UNIDADE') or r.get('ESCOLA') or ''
    cur = r.get('CURSO') or ''
    if c:
        can_c = get_canonical_curso(cur)
        comp = make_composite_key(u, can_c, c)
        index_iftp[comp] += 1

# Aptos Grouping
class_groups = defaultdict(list)
for r in rows_aptos:
    raw_escola = str(r.get('ESCOLA', '')).strip()
    raw_curso = str(r.get('CURSO', '')).strip()
    raw_turma = str(r.get('TURMADESCRICAO', '')).strip()
    unidade_op = get_unidade_operacional(raw_escola, raw_turma)
    can_curso = get_canonical_curso(raw_curso)
    class_groups[(unidade_op, can_curso, raw_turma, raw_escola)].append(r)

print(f"Total Unique Class Groups in Aptos: {len(class_groups)}")

# Macro output rows
macro_rows = []
neg_inaptos_cases = []
dr_issues = []
ited_unidade_issues = []

for (unidade_op, can_curso, turma_desc, raw_escola), recs in class_groups.items():
    dr = 'SENAI-BA'
    if dr != 'SENAI-BA':
        dr_issues.append((unidade_op, turma_desc))
        
    modalidade = get_modalidade(raw_escola)
    turma_code = extract_turma_code(turma_desc)
    turma_display = get_turma_display(turma_desc)
    aptos_count = len(recs)
    
    if is_senai_ited(raw_escola) and unidade_op.lower() == 'senai ited':
        ited_unidade_issues.append((raw_escola, turma_desc, unidade_op))
        
    comp_key = make_composite_key(unidade_op, can_curso, turma_code)
    count_chp = index_chp.get(comp_key, 0)
    count_iftp = index_iftp.get(comp_key, 0)
    
    cats = []
    if count_chp > 0:
        cats.append(('CHP', count_chp))
    if count_iftp > 0:
        cats.append(('IFTP', count_iftp))
    if not cats:
        cats.append(('CHP', 0))
        
    for cat, alunos in cats:
        inaptos = max(0, alunos - aptos_count)
        raw_diff = alunos - aptos_count
        if raw_diff < 0:
            neg_inaptos_cases.append({
                'unidade': unidade_op,
                'curso': can_curso,
                'turma': turma_code,
                'categoria': cat,
                'alunos': alunos,
                'aptos': aptos_count,
                'diff': raw_diff
            })
            
        macro_rows.append({
            'Departamento Regional (DR)': dr,
            'Unidade operacional (Escola)': unidade_op,
            'Curso': can_curso,
            'Categoria': cat,
            'Modalidade': modalidade,
            'Turma': turma_display,
            'Alunos': alunos,
            'Alunos Aptos/Agendados': aptos_count,
            'Alunos Inaptos/Não Agendados': inaptos,
            'Presentes': 0,
            'Presentes (%)': '0.0%',
            'Ausentes': 0,
            'Ausentes (%)': '0.0%'
        })

print(f"\nConsolidated Macro Rows Generated: {len(macro_rows)}")
matched_count = sum(1 for r in macro_rows if r['Alunos'] > 0)
unmatched_count = sum(1 for r in macro_rows if r['Alunos'] == 0)
total_alunos_sum = sum(r['Alunos'] for r in macro_rows)
total_aptos_sum = sum(r['Alunos Aptos/Agendados'] for r in macro_rows)
total_inaptos_sum = sum(r['Alunos Inaptos/Não Agendados'] for r in macro_rows)

print(f"  - Matched Classes: {matched_count} (out of {len(macro_rows)})")
print(f"  - Unmatched Classes: {unmatched_count}")
print(f"  - Total Alunos (Matriculados): {total_alunos_sum}")
print(f"  - Total Aptos/Agendados: {total_aptos_sum}")
print(f"  - Total Inaptos/Não Agendados: {total_inaptos_sum}")

print(f"\nDR Audit: {len(dr_issues)} issues found.")
print(f"ITED Unidade Extraction Audit: {len(ited_unidade_issues)} issues found.")
print(f"Negative Inaptos Audit (Aptos > Matriculados): {len(neg_inaptos_cases)} occurrences.")
if neg_inaptos_cases:
    for c in neg_inaptos_cases:
        print(f"  * Unidade='{c['unidade']}', Turma='{c['turma']}', Categoria={c['categoria']}: Matriculados={c['alunos']} < Aptos={c['aptos']} (Diff: {c['diff']})")

# 3. Check specific confirmed classes
print("\n=================================================================")
print("=== 3. AUDITING CONFIRMED KEY TURMAS ===")
print("=================================================================")

t_96225 = [r for r in macro_rows if r['Turma'] == '96225']
print(f"Turma 96225: {len(t_96225)} row(s)")
for r in t_96225:
    print(f"  - Unidade: '{r['Unidade operacional (Escola)']}', Curso: '{r['Curso']}', Categ: '{r['Categoria']}', Alunos: {r['Alunos']}, Aptos: {r['Alunos Aptos/Agendados']}, Inaptos: {r['Alunos Inaptos/Não Agendados']}")

t_96226 = [r for r in macro_rows if r['Turma'] == '96226']
print(f"Turma 96226: {len(t_96226)} row(s)")
for r in t_96226:
    print(f"  - Unidade: '{r['Unidade operacional (Escola)']}', Curso: '{r['Curso']}', Categ: '{r['Categoria']}', Alunos: {r['Alunos']}, Aptos: {r['Alunos Aptos/Agendados']}, Inaptos: {r['Alunos Inaptos/Não Agendados']}")

t_93629 = [r for r in macro_rows if r['Turma'] == '93629']
print(f"Turma 93629: {len(t_93629)} row(s)")
for r in t_93629:
    print(f"  - Unidade: '{r['Unidade operacional (Escola)']}', Curso: '{r['Curso']}', Categ: '{r['Categoria']}', Alunos: {r['Alunos']}, Aptos: {r['Alunos Aptos/Agendados']}, Inaptos: {r['Alunos Inaptos/Não Agendados']}")

# 4. Multi-Unit Class Audit (e.g. 20251CHPELET7)
t_elet7 = [r for r in macro_rows if r['Turma'] == '20251CHPELET7']
print(f"\nMulti-Unit Turma 20251CHPELET7: {len(t_elet7)} distinct rows:")
for r in t_elet7:
    print(f"  - Unidade: '{r['Unidade operacional (Escola)']}', Alunos={r['Alunos']}, Aptos={r['Alunos Aptos/Agendados']}")

# 5. IFTP EMB/EMP Suffix Aggregation Audit (e.g. 97360, 92173)
t_97360 = [r for r in macro_rows if r['Turma'] == '97360']
print(f"\nIFTP Turma 97360: {len(t_97360)} row(s):")
for r in t_97360:
    print(f"  - Unidade: '{r['Unidade operacional (Escola)']}', Curso: '{r['Curso']}', Categ: '{r['Categoria']}', Alunos={r['Alunos']}, Aptos={r['Alunos Aptos/Agendados']}")
