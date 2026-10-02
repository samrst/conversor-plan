import openpyxl
import xlrd
import unicodedata
import re
from collections import defaultdict

def normalize_string(s):
    if not s: return ''
    nfd = unicodedata.normalize('NFD', str(s))
    no_accents = ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]', '', no_accents.lower())

def normalize_cod_turma(cod):
    if not cod: return ''
    s = str(cod).strip()
    if not s: return ''
    s = re.sub(r'^(?:saep|[tbg])[\s_-]*', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'[\s_-]*(?:emb|emp)$', '', s, flags=re.IGNORECASE).strip()
    return s

def extract_turma_code(turma_desc):
    if not turma_desc: return ''
    s = str(turma_desc).strip()
    if not s: return ''
    raw_code = s
    for sep in ['-', '–', '—']:
        if sep in raw_code:
            raw_code = raw_code.split(sep)[0].strip()
            break
    return normalize_cod_turma(raw_code)

def is_senai_ited(escola):
    if not escola: return False
    norm = normalize_string(escola)
    return 'senaiited' in norm or norm == 'ited'

def get_unidade_operacional(escola, turma_desc):
    raw_escola = (escola or '').strip()
    if is_senai_ited(raw_escola):
        desc = (turma_desc or '').strip()
        for sep in ['-', '–', '—']:
            if sep in desc:
                parts = desc.split(sep)
                extracted = sep.join(parts[1:]).strip()
                if extracted:
                    return extracted
        return desc or raw_escola
    return raw_escola

def normalize_unidade(unidade):
    if not unidade: return ''
    s = normalize_string(unidade)
    s = re.sub(r'^(?:senai|cfp|escola|unidade|centro|cetind)+', '', s)
    canonical = {
        'feiradesantana': 'feira',
        'feiradesantanabest': 'feira',
        'feiraserrinha': 'feira',
        'vitoriadaconquista': 'conquista',
        'luiseduardomagalhaes': 'lem',
        'luiseduardo': 'lem',
        'salvadorcimatec': 'cimatec',
        'salvadordendezeiros': 'dendezeiros',
    }
    for k, v in canonical.items():
        if s.startswith(k) or s == k:
            return v
    return s

def normalize_curso(curso):
    if not curso: return ''
    s = normalize_string(curso)
    s = re.sub(r'^(?:tecnicoem|tecnicode|cursobtecnico|curso)+', '', s)
    return s

def make_composite_key(unidade, curso, turma_code):
    u = normalize_unidade(unidade)
    c = normalize_curso(curso)
    t = normalize_cod_turma(turma_code)
    return f"{u}___{c}___{t}"

# 1. Load CHP
wb_chp = xlrd.open_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
ws_chp = wb_chp.sheet_by_name('Dados')
index_chp_comp = defaultdict(int)
for r in range(2, ws_chp.nrows):
    unid = ws_chp.cell_value(r, 0)
    curso = ws_chp.cell_value(r, 2)
    cod = ws_chp.cell_value(r, 4)
    if cod:
        key = make_composite_key(unid, curso, cod)
        index_chp_comp[key] += 1

# 2. Load IFTP
wb_iftp = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx', data_only=True)
ws_iftp = wb_iftp['Dados']
index_iftp_comp = defaultdict(int)
for r in range(2, ws_iftp.max_row + 1):
    unid = ws_iftp.cell(r, 1).value
    curso = ws_iftp.cell(r, 8).value
    cod = ws_iftp.cell(r, 9).value
    if cod:
        key = make_composite_key(unid, curso, cod)
        index_iftp_comp[key] += 1

# 3. Load Aptos
wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
ws_aptos = wb_aptos['Estudante']
class_groups = defaultdict(list)
for r in range(2, ws_aptos.max_row + 1):
    esc = ws_aptos.cell(r, 5).value
    desc = ws_aptos.cell(r, 6).value
    curso = ws_aptos.cell(r, 19).value
    
    unidade_op = get_unidade_operacional(esc, desc)
    turma_code = extract_turma_code(desc)
    
    group_key = (esc, unidade_op, curso, desc, turma_code)
    class_groups[group_key].append(r)

print(f"Total turmas no índice CHP composto: {len(index_chp_comp)}")
print(f"Total turmas no índice IFTP composto: {len(index_iftp_comp)}")
print(f"Total turmas agrupadas em Aptos: {len(class_groups)}")

results = []
matched_count = 0
unmatched_count = 0

for (esc, unidade_op, curso, desc, turma_code), recs in class_groups.items():
    aptos_count = len(recs)
    comp_key = make_composite_key(unidade_op, curso, turma_code)
    
    count_chp = index_chp_comp.get(comp_key, 0)
    count_iftp = index_iftp_comp.get(comp_key, 0)
    
    cats = []
    if count_chp > 0:
        cats.append(('CHP', count_chp))
    if count_iftp > 0:
        cats.append(('IFTP', count_iftp))
        
    if not cats:
        cats.append(('CHP', 0))
        unmatched_count += 1
    else:
        matched_count += 1
        
    for cat, alunos in cats:
        results.append({
            'Unidade': unidade_op,
            'Curso': curso,
            'Turma': turma_code,
            'Categoria': cat,
            'Alunos': alunos,
            'Aptos': aptos_count,
            'Inaptos': max(0, alunos - aptos_count)
        })

print(f"\nResumo da Consolidação:")
print(f"Turmas Casadas com Alunos > 0: {matched_count}")
print(f"Turmas Sem Correspondência (Alunos = 0): {unmatched_count}")
print(f"Total Alunos: {sum(r['Alunos'] for r in results)}")
print(f"Total Aptos: {sum(r['Aptos'] for r in results)}")

print("\n--- AMOSTRA DE TURMAS CONSOLIDADAS COM CHAVE COMPOSTA ---")
for r in results[:25]:
    print(f"Unidade: {r['Unidade'][:20]:<20} | Curso: {r['Curso'][:25]:<25} | Turma: {r['Turma']:<15} | Cat: {r['Categoria']:<5} | Alunos: {r['Alunos']:<3} | Aptos: {r['Aptos']:<3} | Inaptos: {r['Inaptos']:<3}")
