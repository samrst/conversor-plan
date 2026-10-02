import openpyxl
import xlrd
import os
import re

print("="*80)
print("SIMULAÇÃO DETALHADA DO FLUXO COMPLETO E AUDITORIA DOS DADOS")
print("="*80)

def normalize_turma_code(val):
    if not val:
        return ''
    s = str(val).strip()
    s = re.sub(r'^(?:SAEP|[TBG])', '', s, flags=re.IGNORECASE)
    return s.strip()

def extract_turma_code_aptos(val):
    if not val:
        return ''
    s = str(val).strip()
    # Match code before dash or space
    # e.g. "G20251CHPELET7 - DENDEZEIROS" -> "G20251CHPELET7"
    # "SAEP96377" -> "SAEP96377"
    m = re.match(r'^([^\s\-]+)', s)
    if m:
        code = m.group(1).strip()
    else:
        code = s
    return code

# 1. IFTP analysis
iftp_file = r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx'
wb_iftp = openpyxl.load_workbook(iftp_file, data_only=True)

print("\n--- DETALHES DO ARQUIVO IFTP ---")
ws_dados_iftp = wb_iftp['Dados']
headers_iftp = [ws_dados_iftp.cell(1, c).value for c in range(1, ws_dados_iftp.max_column + 1)]
print("Cabeçalhos da aba 'Dados':", headers_iftp)
codturma_col_iftp = headers_iftp.index('CODTURMA') + 1

sample_iftp = []
for r in range(2, 17):
    cell = ws_dados_iftp.cell(r, codturma_col_iftp)
    sample_iftp.append((cell.value, type(cell.value).__name__, cell.data_type))

print("\nAmostra de 15 valores de CODTURMA no IFTP (aba 'Dados'):")
for val, pytype, dtype in sample_iftp:
    norm = normalize_turma_code(val)
    print(f"  Bruto: {repr(val):<20} | Tipo: {pytype:<10} | dtype: {dtype:<5} | Normalizado: {repr(norm)}")

# 2. CHP analysis
chp_file = r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls'
wb_chp = xlrd.open_workbook(chp_file)
print("\n--- DETALHES DO ARQUIVO CHP ---")
ws_dados_chp = wb_chp.sheet_by_name('Dados')
print("Linha 1 da aba 'Dados':", ws_dados_chp.row_values(0)[:10])
print("Linha 2 da aba 'Dados' (Cabeçalho real):", ws_dados_chp.row_values(1)[:10])

headers_chp = ws_dados_chp.row_values(1)
codturma_col_chp = headers_chp.index('CODTURMA')

sample_chp = []
for r in range(2, 17):
    val = ws_dados_chp.cell_value(r, codturma_col_chp)
    ctype = ws_dados_chp.cell_type(r, codturma_col_chp)
    sample_chp.append((val, type(val).__name__, ctype))

print("\nAmostra de 15 valores de CODTURMA no CHP (aba 'Dados'):")
for val, pytype, ctype in sample_chp:
    norm = normalize_turma_code(val)
    print(f"  Bruto: {repr(val):<20} | Tipo: {pytype:<10} | ctype: {ctype:<5} | Normalizado: {repr(norm)}")

# 3. APTOS analysis
aptos_file = r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx'
wb_aptos = openpyxl.load_workbook(aptos_file, data_only=True)
print("\n--- DETALHES DO ARQUIVO APTOS ---")
ws_estudante = wb_aptos['Estudante']
headers_aptos = [ws_estudante.cell(1, c).value for c in range(1, ws_estudante.max_column + 1)]
print("Cabeçalhos da aba 'Estudante':", headers_aptos)
turma_col_aptos = headers_aptos.index('TURMADESCRICAO') + 1

sample_aptos = []
for r in range(2, 25):
    cell = ws_estudante.cell(r, turma_col_aptos)
    sample_aptos.append((cell.value, type(cell.value).__name__))

print("\nAmostra de valores de TURMADESCRICAO em Aptos (aba 'Estudante'):")
unique_turmas_aptos = []
for val, pytype in sample_aptos:
    if val not in unique_turmas_aptos:
        unique_turmas_aptos.append(val)

for val in unique_turmas_aptos[:10]:
    extracted = extract_turma_code_aptos(val)
    norm = normalize_turma_code(extracted)
    print(f"  TURMADESCRICAO: {repr(val):<45} -> Extraído: {repr(extracted):<20} -> Normalizado: {repr(norm)}")
