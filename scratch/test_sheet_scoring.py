import openpyxl
import xlrd
import re
import os

print("="*80)
print("TESTE DE DETECÇÃO DE ABAS E CABEÇALHOS")
print("="*80)

def detect_sheet_header_and_rows(raw_rows):
    # Find best header row among first 10 rows
    best_row_idx = 0
    best_score = -1
    best_headers = []
    
    for r_idx in range(min(10, len(raw_rows))):
        row = raw_rows[r_idx]
        # Count non-empty text strings
        non_empty = [str(cell).strip() for cell in row if cell is not None and str(cell).strip() != '']
        if not non_empty:
            continue
        
        # Calculate score based on column count and known keywords
        keywords = ['codturma', 'turmadescricao', 'turma', 'escola', 'curso', 'unidade', 'cpf', 'nome', 'presentes', 'ausentes']
        kw_matches = sum(1 for c in non_empty if any(k in c.lower() for k in keywords))
        score = len(non_empty) + (kw_matches * 10)
        
        if score > best_score:
            best_score = score
            best_row_idx = r_idx
            best_headers = [str(c).strip() if c is not None else '' for c in row]
            # Trim trailing empties
            while best_headers and best_headers[-1] == '':
                best_headers.pop()

    # Data rows are all rows after best_row_idx
    data_rows = []
    for r_idx in range(best_row_idx + 1, len(raw_rows)):
        row = raw_rows[r_idx]
        obj = {}
        for c_idx, h in enumerate(best_headers):
            if h:
                val = row[c_idx] if c_idx < len(row) else ''
                obj[h] = val
        data_rows.append(obj)
        
    return best_row_idx, best_headers, data_rows

def score_sheet(sheet_name, headers, row_count, role=None):
    score = 0
    clean_name = sheet_name.lower()
    
    # Sheet name bonuses/penalties
    if any(w in clean_name for w in ['dados', 'estudante', 'aluno', 'microdado', 'base']):
        score += 50
    if any(w in clean_name for w in ['resumo', 'sumario', 'pivot', 'filtro']):
        score -= 50
        
    # Role-based keywords
    role_keywords = {
        'aptos': ['turmadescricao', 'escola', 'curso', 'cpf', 'aluno', 'estudante'],
        'matriculados_chp': ['codturma', 'unidade', 'curso', 'matriculado'],
        'matriculados_iftp': ['codturma', 'unidade', 'curso', 'matriculado'],
        'participacao': ['presentes', 'ausentes', 'turma', 'escola', 'curso']
    }
    
    target_kws = role_keywords.get(role, ['codturma', 'turmadescricao', 'escola', 'curso'])
    
    clean_headers = [h.lower() for h in headers]
    for kw in target_kws:
        if any(kw in h for h in clean_headers):
            score += 30
            
    # Row count weight (more data rows is almost always the microdata sheet)
    score += min(50, row_count // 10)
    return score

# Test with real files
print("\n1. Testando Aptos (SAEP 2026-2_28_08_2026_BA.xlsx):")
wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
for sname in wb_aptos.sheetnames:
    ws = wb_aptos[sname]
    raw_rows = []
    for r in range(1, ws.max_row + 1):
        raw_rows.append([ws.cell(r, c).value for c in range(1, ws.max_column + 1)])
    h_idx, headers, rows = detect_sheet_header_and_rows(raw_rows)
    s = score_sheet(sname, headers, len(rows), 'aptos')
    print(f"  Aba: '{sname:<25}' | Header Row: {h_idx+1} | Cols: {len(headers)} | Rows: {len(rows):<5} | Score: {s}")

print("\n2. Testando CHP (Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls):")
wb_chp = xlrd.open_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
for sname in wb_chp.sheet_names():
    ws = wb_chp.sheet_by_name(sname)
    raw_rows = [ws.row_values(r) for r in range(ws.nrows)]
    h_idx, headers, rows = detect_sheet_header_and_rows(raw_rows)
    s = score_sheet(sname, headers, len(rows), 'matriculados_chp')
    print(f"  Aba: '{sname:<25}' | Header Row: {h_idx+1} | Cols: {len(headers)} | Rows: {len(rows):<5} | Score: {s}")

print("\n3. Testando IFTP (Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx):")
wb_iftp = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx', data_only=True)
for sname in wb_iftp.sheetnames:
    ws = wb_iftp[sname]
    raw_rows = []
    for r in range(1, ws.max_row + 1):
        raw_rows.append([ws.cell(r, c).value for c in range(1, ws.max_column + 1)])
    h_idx, headers, rows = detect_sheet_header_and_rows(raw_rows)
    s = score_sheet(sname, headers, len(rows), 'matriculados_iftp')
    print(f"  Aba: '{sname:<25}' | Header Row: {h_idx+1} | Cols: {len(headers)} | Rows: {len(rows):<5} | Score: {s}")
