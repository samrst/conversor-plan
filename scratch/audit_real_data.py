import openpyxl
import xlrd
import os
import re

print("="*80)
print("AUDITORIA DE DADOS REAIS - PROVA OBJETIVA")
print("="*80)

# 1. Inspect IFTP
iftp_file = r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx'
wb_iftp = openpyxl.load_workbook(iftp_file, data_only=True)
print("\n--- 1. ARQUIVO IFTP:", os.path.basename(iftp_file), "---")
print("Abas encontradas:", wb_iftp.sheetnames)
for sname in wb_iftp.sheetnames:
    ws = wb_iftp[sname]
    print(f"\n[Aba: '{sname}'] - Linhas: {ws.max_row}, Colunas: {ws.max_column}")
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    print(f"  Linha 1 (cabeçalhos): {headers[:10]}")
    if ws.max_row > 1:
        row2 = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        print(f"  Linha 2 (primeira linha de dados): {row2[:10]}")

# 2. Inspect CHP (.xls)
chp_file = r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls'
wb_chp = xlrd.open_workbook(chp_file)
print("\n--- 2. ARQUIVO CHP:", os.path.basename(chp_file), "---")
print("Abas encontradas:", wb_chp.sheet_names())
for sname in wb_chp.sheet_names():
    ws = wb_chp.sheet_by_name(sname)
    print(f"\n[Aba: '{sname}'] - Linhas: {ws.nrows}, Colunas: {ws.ncols}")
    for r in range(min(5, ws.nrows)):
        print(f"  Linha {r+1}: {ws.row_values(r)[:10]}")

# 3. Inspect Aptos
aptos_file = r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx'
wb_aptos = openpyxl.load_workbook(aptos_file, data_only=True)
print("\n--- 3. ARQUIVO APTOS:", os.path.basename(aptos_file), "---")
print("Abas encontradas:", wb_aptos.sheetnames)
for sname in wb_aptos.sheetnames:
    ws = wb_aptos[sname]
    if ws.max_row > 1:
        headers = [ws.cell(1, c).value for c in range(1, min(15, ws.max_column + 1))]
        print(f"[Aba: '{sname}'] - Linhas: {ws.max_row}, Colunas: {ws.max_column} | Cabeçalhos: {headers[:6]}")

# 4. Inspect Participacao
part_file = r'C:\Users\sam.sampaio\Downloads\RelatorioParticipacaoTurmasProvas_2026-06-22_14-42-12.xlsx'
wb_part = openpyxl.load_workbook(part_file, data_only=True)
print("\n--- 4. ARQUIVO PARTICIPACAO:", os.path.basename(part_file), "---")
print("Abas encontradas:", wb_part.sheetnames)
for sname in wb_part.sheetnames:
    ws = wb_part[sname]
    print(f"[Aba: '{sname}'] - Linhas: {ws.max_row}, Colunas: {ws.max_column}")
    if ws.max_row > 0:
        headers = [ws.cell(1, c).value for c in range(1, min(15, ws.max_column + 1))]
        print(f"  Linha 1: {headers[:10]}")
