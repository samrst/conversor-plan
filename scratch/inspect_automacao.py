import openpyxl
import xlrd
from collections import defaultdict

wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
ws_aptos = wb_aptos['Estudante']
headers_aptos = [ws_aptos.cell(1, c).value for c in range(1, ws_aptos.max_column + 1)]
rows_aptos = []
for r in range(2, ws_aptos.max_row + 1):
    vals = [ws_aptos.cell(r, c).value for c in range(1, ws_aptos.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_aptos.append(dict(zip(headers_aptos, vals)))

auto_aptos = defaultdict(list)
for r in rows_aptos:
    cur = str(r.get('CURSO', ''))
    if 'automa' in cur.lower():
        auto_aptos[(r.get('ESCOLA'), cur, r.get('TURMADESCRICAO'))].append(r)

print("=== AUTOMACAO TURMAS IN APTOS ===")
for (esc, cur, tdesc), recs in auto_aptos.items():
    print(f"Escola: '{esc}' | Curso: '{cur}' | TurmaDesc: '{tdesc}' -> {len(recs)} aptos")
