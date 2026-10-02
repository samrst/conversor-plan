import openpyxl
import xlrd
import os
import re
import unicodedata
from collections import defaultdict

def normalize_string(s):
    if not s: return ''
    nfd = unicodedata.normalize('NFD', str(s))
    no_accents = ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]', '', no_accents.lower())

def normalize_cod_turma(cod):
    if not cod: return ''
    s = str(cod).strip()
    if not s: return ''
    s = re.sub(r'^(?:saep|[tbg])[\s_-]*', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'[\s_-]*(?:emb|emp)$', '', s, flags=re.IGNORECASE).strip()
    return s

def extract_turma_code(turma_desc):
    if not turma_desc: return ''
    s = str(turma_desc).strip()
    if not s: return ''
    raw_code = s
    for sep in ['-', '–', '—']:
        if sep in raw_code:
            raw_code = raw_code.split(sep)[0].strip()
            break
    return normalize_cod_turma(raw_code)

def is_senai_ited(escola):
    if not escola: return False
    norm = normalize_string(escola)
    return 'senaiited' in norm or norm == 'ited'

def get_unidade_operacional(escola, turma_desc):
    raw_escola = (escola or '').strip()
    if is_senai_ited(raw_escola):
        desc = (turma_desc or '').strip()
        for sep in ['-', '–', '—']:
            if sep in desc:
                parts = desc.split(sep)
                extracted = sep.join(parts[1:]).strip()
                if extracted:
                    return extracted
        return desc or raw_escola
    return raw_escola

def normalize_unidade(unidade):
    if not unidade: return ''
    s = normalize_string(unidade)
    s = re.sub(r'^(?:senai|cfp|escola|unidade|centro|cetind)+', '', s)
    canonical = {
        'feiradesantana': 'feira',
        'feiradesantanabest': 'feira',
        'feiraserrinha': 'feira',
        'vitoriadaconquista': 'conquista',
        'luiseduardomagalhaes': 'lem',
        'luiseduardo': 'lem',
        'salvadorcimatec': 'cimatec',
        'salvadordendezeiros': 'dendezeiros'
    }
    for k, v in canonical.items():
        if s.startswith(k) or s == k:
            return v
    return s

def normalize_curso(curso):
    if not curso: return ''
    s = normalize_string(curso)
    s = re.sub(r'^(?:tecnicoem|tecnicode|cursobtecnico|curso)+', '', s)
    return s

def make_composite_key(unidade, curso, turma_code):
    u = normalize_unidade(unidade)
    c = normalize_curso(curso)
    t = normalize_cod_turma(turma_code)
    return f"{u}___{c}___{t}"

# 1. Load Aptos
print("Loading Aptos...")
wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
ws_aptos = wb_aptos['Estudante']
headers_aptos = [ws_aptos.cell(1, c).value for c in range(1, ws_aptos.max_column + 1)]
rows_aptos = []
for r in range(2, ws_aptos.max_row + 1):
    vals = [ws_aptos.cell(r, c).value for c in range(1, ws_aptos.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_aptos.append(dict(zip(headers_aptos, vals)))

# 2. Load CHP
print("Loading CHP...")
wb_chp = xlrd.open_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
ws_chp = wb_chp.sheet_by_name('Dados')
headers_chp = [ws_chp.cell_value(1, c) for c in range(ws_chp.ncols)]
rows_chp = []
for r in range(2, ws_chp.nrows):
    vals = [ws_chp.cell_value(r, c) for c in range(ws_chp.ncols)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_chp.append(dict(zip(headers_chp, vals)))

# 3. Load IFTP
print("Loading IFTP...")
wb_iftp = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx', data_only=True)
ws_iftp = wb_iftp['Dados']
headers_iftp = [ws_iftp.cell(1, c).value for c in range(1, ws_iftp.max_column + 1)]
rows_iftp = []
for r in range(2, ws_iftp.max_row + 1):
    vals = [ws_iftp.cell(r, c).value for c in range(1, ws_iftp.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_iftp.append(dict(zip(headers_iftp, vals)))

print(f"Total rows in Aptos Estudante: {len(rows_aptos)}")
print(f"Total rows in CHP Dados: {len(rows_chp)}")
print(f"Total rows in IFTP Dados: {len(rows_iftp)}")

# Detailed Indexing
index_chp_comp = defaultdict(int)
index_chp_turma_only = defaultdict(int)
for r in rows_chp:
    c = r.get('CODTURMA')
    u = r.get('UNIDADE') or r.get('ESCOLA') or ''
    cur = r.get('CURSO') or ''
    if c:
        comp = make_composite_key(u, cur, c)
        norm_t = normalize_cod_turma(c)
        index_chp_comp[comp] += 1
        index_chp_turma_only[norm_t] += 1

index_iftp_comp = defaultdict(int)
index_iftp_turma_only = defaultdict(int)
for r in rows_iftp:
    c = r.get('CODTURMA')
    u = r.get('UNIDADE') or r.get('ESCOLA') or ''
    cur = r.get('CURSO') or ''
    if c:
        comp = make_composite_key(u, cur, c)
        norm_t = normalize_cod_turma(c)
        index_iftp_comp[comp] += 1
        index_iftp_turma_only[norm_t] += 1

# Aptos groups
class_groups = defaultdict(list)
for r in rows_aptos:
    raw_escola = str(r.get('ESCOLA', '')).strip()
    raw_curso = str(r.get('CURSO', '')).strip()
    raw_turma = str(r.get('TURMADESCRICAO', '')).strip()
    unidade_op = get_unidade_operacional(raw_escola, raw_turma)
    class_groups[(unidade_op, raw_curso, raw_turma, raw_escola)].append(r)

print(f"\nTotal distinct class groups in Aptos: {len(class_groups)}")

# Audit each of the class groups
matches = []
unmatched = []
old_total_alunos = 0
new_total_alunos = 0

for (unidade_op, curso, turma_desc, escola_raw), recs in class_groups.items():
    turma_code = extract_turma_code(turma_desc)
    aptos_count = len(recs)
    comp_key = make_composite_key(unidade_op, curso, turma_code)
    
    # Old logic count:
    old_chp = index_chp_turma_only.get(turma_code, 0)
    old_iftp = index_iftp_turma_only.get(turma_code, 0)
    old_alunos = old_chp + old_iftp
    old_total_alunos += old_alunos
    
    # New logic count:
    chp_count = index_chp_comp.get(comp_key, 0)
    iftp_count = index_iftp_comp.get(comp_key, 0)
    
    cat = 'CHP'
    alunos = 0
    if chp_count > 0 and iftp_count > 0:
        cat = 'CHP+IFTP'
        alunos = chp_count + iftp_count
    elif chp_count > 0:
        cat = 'CHP'
        alunos = chp_count
    elif iftp_count > 0:
        cat = 'IFTP'
        alunos = iftp_count
    else:
        cat = 'CHP (0)'
        alunos = 0
        
    new_total_alunos += alunos
    
    status = 'MATCH' if alunos > 0 else 'SEM MATCH'
    
    info = {
        'unidade': unidade_op,
        'curso': curso,
        'categoria': cat,
        'turma': turma_code,
        'turma_desc': turma_desc,
        'aptos': aptos_count,
        'alunos': alunos,
        'old_alunos': old_alunos,
        'comp_key': comp_key,
        'status': status
    }
    
    if status == 'MATCH':
        matches.append(info)
    else:
        unmatched.append(info)

print(f"\n=== MATCH SUMMARY ===")
print(f"Total groups: {len(class_groups)}")
print(f"Matches: {len(matches)}")
print(f"Unmatched (SEM MATCH): {len(unmatched)}")
print(f"Old total alunos: {old_total_alunos}")
print(f"New total alunos: {new_total_alunos}")

print(f"\n=== UNMATCHED (SEM MATCH) DETAILS ===")
for u in unmatched:
    print(f"Unidade: {u['unidade']}")
    print(f"Curso: {u['curso']}")
    print(f"Turma Code: {u['turma']}")
    print(f"TURMADESCRICAO: {u['turma_desc']}")
    print(f"Aptos count: {u['aptos']}")
    print(f"Composite Key: {u['comp_key']}")
    print(f"Old Alunos (Turma only): {u['old_alunos']}")
    print("-" * 50)
