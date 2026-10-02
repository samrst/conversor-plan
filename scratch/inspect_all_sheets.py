import os
import openpyxl

downloads_dir = r"C:\Users\sam.sampaio\Downloads"

def analyze_all_sheets(filename):
    path = os.path.join(downloads_dir, filename)
    print(f"\n=================================================================")
    print(f"FILE: {filename}")
    print(f"=================================================================")
    wb = openpyxl.load_workbook(path, data_only=True)
    print("All Sheet Names:", wb.sheetnames)
    for idx, sname in enumerate(wb.sheetnames):
        ws = wb[sname]
        print(f"\n--- Sheet Index {idx}: '{sname}' (rows={ws.max_row}, cols={ws.max_column}) ---")
        for r in range(1, min(6, ws.max_row + 1)):
            row_vals = [ws.cell(r, c).value for c in range(1, min(20, ws.max_column + 1))]
            print(f"  Row {r:2d}: {row_vals}")

if __name__ == '__main__':
    analyze_all_sheets("Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx")
    analyze_all_sheets("SAEP 2026-2_28_08_2026_BA.xlsx")
