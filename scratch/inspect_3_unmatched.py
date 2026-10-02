import openpyxl
import xlrd
import re
import unicodedata

# 1. Load Aptos
wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
ws_aptos = wb_aptos['Estudante']
headers_aptos = [ws_aptos.cell(1, c).value for c in range(1, ws_aptos.max_column + 1)]
rows_aptos = []
for r in range(2, ws_aptos.max_row + 1):
    vals = [ws_aptos.cell(r, c).value for c in range(1, ws_aptos.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_aptos.append(dict(zip(headers_aptos, vals)))

# 2. Load CHP
wb_chp = xlrd.open_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
ws_chp = wb_chp.sheet_by_name('Dados')
headers_chp = [ws_chp.cell_value(1, c) for c in range(ws_chp.ncols)]
rows_chp = []
for r in range(2, ws_chp.nrows):
    vals = [ws_chp.cell_value(r, c) for c in range(ws_chp.ncols)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_chp.append(dict(zip(headers_chp, vals)))

# 3. Load IFTP
wb_iftp = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx', data_only=True)
ws_iftp = wb_iftp['Dados']
headers_iftp = [ws_iftp.cell(1, c).value for c in range(1, ws_iftp.max_column + 1)]
rows_iftp = []
for r in range(2, ws_iftp.max_row + 1):
    vals = [ws_iftp.cell(r, c).value for c in range(1, ws_iftp.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_iftp.append(dict(zip(headers_iftp, vals)))

print("=== EXACT DETAILS FOR 96225 ===")
for r in rows_aptos:
    if '96225' in str(r.get('TURMADESCRICAO')):
        print(f"APTOS: Escola='{r.get('ESCOLA')}', Curso='{r.get('CURSO')}', TurmaDesc='{r.get('TURMADESCRICAO')}'")
        break
for r in rows_chp:
    if '96225' in str(r.get('CODTURMA')):
        print(f"CHP: Unidade='{r.get('UNIDADE') or r.get('ESCOLA')}', Curso='{r.get('CURSO')}', CodTurma='{r.get('CODTURMA')}'")
for r in rows_iftp:
    if '96225' in str(r.get('CODTURMA')):
        print(f"IFTP: Unidade='{r.get('UNIDADE') or r.get('ESCOLA')}', Curso='{r.get('CURSO')}', CodTurma='{r.get('CODTURMA')}'")

print("\n=== EXACT DETAILS FOR 96226 ===")
for r in rows_aptos:
    if '96226' in str(r.get('TURMADESCRICAO')):
        print(f"APTOS: Escola='{r.get('ESCOLA')}', Curso='{r.get('CURSO')}', TurmaDesc='{r.get('TURMADESCRICAO')}'")
        break
for r in rows_chp:
    if '96226' in str(r.get('CODTURMA')):
        print(f"CHP: Unidade='{r.get('UNIDADE') or r.get('ESCOLA')}', Curso='{r.get('CURSO')}', CodTurma='{r.get('CODTURMA')}'")
for r in rows_iftp:
    if '96226' in str(r.get('CODTURMA')):
        print(f"IFTP: Unidade='{r.get('UNIDADE') or r.get('ESCOLA')}', Curso='{r.get('CURSO')}', CodTurma='{r.get('CODTURMA')}'")

print("\n=== EXACT DETAILS FOR 93629 ===")
for r in rows_aptos:
    if '93629' in str(r.get('TURMADESCRICAO')):
        print(f"APTOS: Escola='{r.get('ESCOLA')}', Curso='{r.get('CURSO')}', TurmaDesc='{r.get('TURMADESCRICAO')}'")
        break
for r in rows_chp:
    if '93629' in str(r.get('CODTURMA')):
        print(f"CHP: Unidade='{r.get('UNIDADE') or r.get('ESCOLA')}', Curso='{r.get('CURSO')}', CodTurma='{r.get('CODTURMA')}'")
for r in rows_iftp:
    if '93629' in str(r.get('CODTURMA')):
        print(f"IFTP: Unidade='{r.get('UNIDADE') or r.get('ESCOLA')}', Curso='{r.get('CURSO')}', CodTurma='{r.get('CODTURMA')}'")
