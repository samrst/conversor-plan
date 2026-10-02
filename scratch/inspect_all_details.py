import openpyxl
import xlrd
import os

print("="*80)
print("AUDITORIA COMPLETA DE TODOS OS ARQUIVOS E SUAS ABAS")
print("="*80)

def inspect_xlsx(filepath):
    print(f"\n>>> Arquivo: {os.path.basename(filepath)}")
    wb = openpyxl.load_workbook(filepath, data_only=True)
    print(f"Total de Abas: {len(wb.sheetnames)} -> {wb.sheetnames}")
    for idx, sname in enumerate(wb.sheetnames):
        ws = wb[sname]
        print(f"  Aba [{idx}] '{sname}': {ws.max_row} linhas x {ws.max_column} colunas")
        for r in range(1, min(5, ws.max_row + 1)):
            row_vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
            if any(v is not None for v in row_vals):
                print(f"    Linha {r}: {row_vals}")

def inspect_xls(filepath):
    print(f"\n>>> Arquivo: {os.path.basename(filepath)}")
    wb = xlrd.open_workbook(filepath)
    print(f"Total de Abas: {len(wb.sheet_names())} -> {wb.sheet_names()}")
    for idx, sname in enumerate(wb.sheet_names()):
        ws = wb.sheet_by_name(sname)
        print(f"  Aba [{idx}] '{sname}': {ws.nrows} linhas x {ws.ncols} colunas")
        for r in range(min(5, ws.nrows)):
            row_vals = ws.row_values(r)[:12]
            if any(v != '' for v in row_vals):
                print(f"    Linha {r+1}: {row_vals}")

# Check files
inspect_xlsx(r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx')
inspect_xls(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
inspect_xlsx(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx')
inspect_xlsx(r'C:\Users\sam.sampaio\Downloads\RelatorioParticipacaoTurmasProvas_2026-06-22_14-42-12.xlsx')
