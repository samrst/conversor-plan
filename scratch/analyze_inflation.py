import openpyxl
import xlrd
from collections import defaultdict
import unicodedata
import re

def normalize_string(s):
    if not s: return ''
    nfd = unicodedata.normalize('NFD', str(s))
    no_accents = ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]', '', no_accents.lower())

def normalize_cod_turma(cod):
    if not cod: return ''
    s = str(cod).strip()
    if not s: return ''
    s = re.sub(r'^(?:saep|[tbg])[\s_-]*', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'[\s_-]*(?:emb|emp)$', '', s, flags=re.IGNORECASE).strip()
    return s

def extract_turma_code(turma_desc):
    if not turma_desc: return ''
    s = str(turma_desc).strip()
    if not s: return ''
    raw_code = s
    for sep in ['-', '–', '—']:
        if sep in raw_code:
            raw_code = raw_code.split(sep)[0].strip()
            break
    return normalize_cod_turma(raw_code)

def is_senai_ited(escola):
    if not escola: return False
    norm = normalize_string(escola)
    return 'senaiited' in norm or norm == 'ited'

def get_unidade_operacional(escola, turma_desc):
    raw_escola = (escola or '').strip()
    if is_senai_ited(raw_escola):
        desc = (turma_desc or '').strip()
        for sep in ['-', '–', '—']:
            if sep in desc:
                parts = desc.split(sep)
                extracted = sep.join(parts[1:]).strip()
                if extracted:
                    return extracted
        return desc or raw_escola
    return raw_escola

def normalize_unidade(unidade):
    if not unidade: return ''
    s = normalize_string(unidade)
    s = re.sub(r'^(?:senai|cfp|escola|unidade|centro|cetind)+', '', s)
    canonical = {
        'feiradesantana': 'feira',
        'feiradesantanabest': 'feira',
        'feiraserrinha': 'feira',
        'vitoriadaconquista': 'conquista',
        'luiseduardomagalhaes': 'lem',
        'luiseduardo': 'lem',
        'salvadorcimatec': 'cimatec',
        'salvadordendezeiros': 'dendezeiros'
    }
    for k, v in canonical.items():
        if s.startswith(k) or s == k:
            return v
    return s

def normalize_curso(curso):
    if not curso: return ''
    s = normalize_string(curso)
    s = re.sub(r'^(?:tecnicoem|tecnicode|cursobtecnico|curso)+', '', s)
    return s

def make_composite_key(unidade, curso, turma_code):
    u = normalize_unidade(unidade)
    c = normalize_curso(curso)
    t = normalize_cod_turma(turma_code)
    return f"{u}___{c}___{t}"

# 1. Load Aptos
wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
ws_aptos = wb_aptos['Estudante']
headers_aptos = [ws_aptos.cell(1, c).value for c in range(1, ws_aptos.max_column + 1)]
rows_aptos = []
for r in range(2, ws_aptos.max_row + 1):
    vals = [ws_aptos.cell(r, c).value for c in range(1, ws_aptos.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_aptos.append(dict(zip(headers_aptos, vals)))

# 2. Load CHP
wb_chp = xlrd.open_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
ws_chp = wb_chp.sheet_by_name('Dados')
headers_chp = [ws_chp.cell_value(1, c) for c in range(ws_chp.ncols)]
rows_chp = []
for r in range(2, ws_chp.nrows):
    vals = [ws_chp.cell_value(r, c) for c in range(ws_chp.ncols)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_chp.append(dict(zip(headers_chp, vals)))

# 3. Load IFTP
wb_iftp = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx', data_only=True)
ws_iftp = wb_iftp['Dados']
headers_iftp = [ws_iftp.cell(1, c).value for c in range(1, ws_iftp.max_column + 1)]
rows_iftp = []
for r in range(2, ws_iftp.max_row + 1):
    vals = [ws_iftp.cell(r, c).value for c in range(1, ws_iftp.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_iftp.append(dict(zip(headers_iftp, vals)))

# Group Aptos
class_groups = defaultdict(list)
for r in rows_aptos:
    raw_escola = str(r.get('ESCOLA', '')).strip()
    raw_curso = str(r.get('CURSO', '')).strip()
    raw_turma = str(r.get('TURMADESCRICAO', '')).strip()
    unidade_op = get_unidade_operacional(raw_escola, raw_turma)
    class_groups[(unidade_op, raw_curso, raw_turma)].append(r)

# Turma-only index (Old)
old_chp_idx = defaultdict(int)
for r in rows_chp:
    c = r.get('CODTURMA')
    if c:
        old_chp_idx[normalize_cod_turma(c)] += 1

old_iftp_idx = defaultdict(int)
for r in rows_iftp:
    c = r.get('CODTURMA')
    if c:
        old_iftp_idx[normalize_cod_turma(c)] += 1

# Composite index (New)
new_chp_idx = defaultdict(int)
for r in rows_chp:
    c = r.get('CODTURMA')
    u = r.get('UNIDADE') or r.get('ESCOLA') or ''
    cur = r.get('CURSO') or ''
    if c:
        new_chp_idx[make_composite_key(u, cur, c)] += 1

new_iftp_idx = defaultdict(int)
for r in rows_iftp:
    c = r.get('CODTURMA')
    u = r.get('UNIDADE') or r.get('ESCOLA') or ''
    cur = r.get('CURSO') or ''
    if c:
        new_iftp_idx[make_composite_key(u, cur, c)] += 1

# Analyze inflation
inflation_by_turma = defaultdict(lambda: {'occurrences_in_aptos': 0, 'old_chp_count': 0, 'old_iftp_count': 0, 'new_chp_total': 0, 'new_iftp_total': 0, 'units': []})

for (unidade_op, curso, turma_desc), recs in class_groups.items():
    turma_code = extract_turma_code(turma_desc)
    comp_key = make_composite_key(unidade_op, curso, turma_code)
    
    n_chp = new_chp_idx.get(comp_key, 0)
    n_iftp = new_iftp_idx.get(comp_key, 0)
    
    o_chp = old_chp_idx.get(turma_code, 0)
    o_iftp = old_iftp_idx.get(turma_code, 0)
    
    entry = inflation_by_turma[turma_code]
    entry['occurrences_in_aptos'] += 1
    entry['old_chp_count'] = o_chp
    entry['old_iftp_count'] = o_iftp
    entry['new_chp_total'] += n_chp
    entry['new_iftp_total'] += n_iftp
    entry['units'].append((unidade_op, n_chp, n_iftp))

print("=== RECONCILING INFLATION: TOP MULTI-UNIT TURMAS ===")
total_old_sum = 0
total_new_sum = 0
for tcode, data in sorted(inflation_by_turma.items(), key=lambda x: (x[1]['occurrences_in_aptos'] * (x[1]['old_chp_count'] + x[1]['old_iftp_count'])), reverse=True):
    old_contrib = data['occurrences_in_aptos'] * (data['old_chp_count'] + data['old_iftp_count'])
    new_contrib = data['new_chp_total'] + data['new_iftp_total']
    total_old_sum += old_contrib
    total_new_sum += new_contrib
    if data['occurrences_in_aptos'] > 1 or old_contrib != new_contrib:
        diff = old_contrib - new_contrib
        print(f"Turma '{tcode}': {data['occurrences_in_aptos']} units in Aptos | Old single count={data['old_chp_count']+data['old_iftp_count']} -> Old Contrib={old_contrib} | New Contrib={new_contrib} (Diff: +{diff} inflated)")
        for u, nc, ni in data['units']:
            print(f"   -> Unit '{u}': CHP={nc}, IFTP={ni}")

print(f"\nTotal Old Sum over 126 groups: {total_old_sum}")
print(f"Total New Sum over 126 groups: {total_new_sum}")
print(f"Total Difference due to duplicate counting across units + unassociated: {total_old_sum - total_new_sum}")
