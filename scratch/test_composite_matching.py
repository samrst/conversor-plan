import openpyxl
import xlrd
import unicodedata
import re
from collections import defaultdict

def normalize_string(s):
    if not s: return ''
    nfd = unicodedata.normalize('NFD', str(s))
    no_accents = ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]', '', no_accents.lower())

def normalize_cod_turma(cod):
    if not cod: return ''
    s = str(cod).strip()
    if not s: return ''
    s = re.sub(r'^(?:saep|[tbg])[\s_-]*', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'[\s_-]*(?:emb|emp)$', '', s, flags=re.IGNORECASE).strip()
    return s

def extract_turma_code(turma_desc):
    if not turma_desc: return ''
    s = str(turma_desc).strip()
    if not s: return ''
    raw_code = s
    for sep in ['-', '–', '—']:
        if sep in raw_code:
            raw_code = raw_code.split(sep)[0].strip()
            break
    return normalize_cod_turma(raw_code)

def is_senai_ited(escola):
    if not escola: return False
    norm = normalize_string(escola)
    return 'senaiited' in norm or norm == 'ited'

def get_unidade_operacional(escola, turma_desc):
    raw_escola = (escola or '').strip()
    if is_senai_ited(raw_escola):
        desc = (turma_desc or '').strip()
        for sep in ['-', '–', '—']:
            if sep in desc:
                parts = desc.split(sep)
                extracted = sep.join(parts[1:]).strip()
                if extracted:
                    return extracted
        return desc or raw_escola
    return raw_escola

# Composite key builder
def make_composite_key(unidade, curso, turma_code):
    u = normalize_string(unidade)
    c = normalize_string(curso)
    t = normalize_cod_turma(turma_code)
    return f"{u}___{c}___{t}"

# 1. Index CHP with composite key
wb_chp = xlrd.open_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
ws_chp = wb_chp.sheet_by_name('Dados')
index_chp_comp = defaultdict(int)
for r in range(2, ws_chp.nrows):
    unid = ws_chp.cell_value(r, 0)
    curso = ws_chp.cell_value(r, 2)
    cod = ws_chp.cell_value(r, 4)
    if cod:
        key = make_composite_key(unid, curso, cod)
        index_chp_comp[key] += 1

# 2. Index IFTP with composite key
wb_iftp = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx', data_only=True)
ws_iftp = wb_iftp['Dados']
index_iftp_comp = defaultdict(int)
for r in range(2, ws_iftp.max_row + 1):
    unid = ws_iftp.cell(r, 1).value # UNIDADE
    curso = ws_iftp.cell(r, 8).value # CURSO
    cod = ws_iftp.cell(r, 9).value # CODTURMA
    if cod:
        key = make_composite_key(unid, curso, cod)
        index_iftp_comp[key] += 1

# 3. Load Aptos and group by (Unidade Operacional, Curso, TurmaDescricao)
wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
ws_aptos = wb_aptos['Estudante']
class_groups = defaultdict(list)
for r in range(2, ws_aptos.max_row + 1):
    esc = ws_aptos.cell(r, 5).value
    turma_desc = ws_aptos.cell(r, 6).value
    curso = ws_aptos.cell(r, 19).value
    
    unidade_op = get_unidade_operacional(esc, turma_desc)
    turma_code = extract_turma_code(turma_desc)
    
    group_key = (esc, unidade_op, curso, turma_desc, turma_code)
    class_groups[group_key].append(r)

print(f"Total turmas no índice CHP composto: {len(index_chp_comp)}")
print(f"Total turmas no índice IFTP composto: {len(index_iftp_comp)}")
print(f"Total grupos em Aptos: {len(class_groups)}")

# Match and check 20251CHPELET7 across its units
print("\n--- DETALHAMENTO DA TURMA '20251CHPELET7' POR UNIDADE ---")
for (esc, unidade_op, curso, turma_desc, turma_code), recs in class_groups.items():
    if turma_code == '20251CHPELET7':
        comp_key = make_composite_key(unidade_op, curso, turma_code)
        count_chp = index_chp_comp.get(comp_key, 0)
        print(f"Unidade: {unidade_op:<20} | Curso: {curso:<28} | Aptos: {len(recs):<3} | Matriculados CHP: {count_chp:<3} | Key: {comp_key}")

print("\n--- DETALHAMENTO DA TURMA '97360' (IFTP) POR UNIDADE ---")
for (esc, unidade_op, curso, turma_desc, turma_code), recs in class_groups.items():
    if turma_code == '97360':
        comp_key = make_composite_key(unidade_op, curso, turma_code)
        count_iftp = index_iftp_comp.get(comp_key, 0)
        print(f"Unidade: {unidade_op:<20} | Curso: {curso:<28} | Aptos: {len(recs):<3} | Matriculados IFTP: {count_iftp:<3} | Key: {comp_key}")
