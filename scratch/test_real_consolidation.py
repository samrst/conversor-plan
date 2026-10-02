import openpyxl
import xlrd
import re
from collections import defaultdict

def normalize_string(s):
    if not s: return ''
    import unicodedata
    nfd = unicodedata.normalize('NFD', str(s))
    no_accents = ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]', '', no_accents.lower())

def normalize_cod_turma(cod):
    if not cod: return ''
    s = str(cod).strip()
    if not s: return ''
    # 1. Remove prefix T, B, G, SAEP (case-insensitive) only at the start
    s = re.sub(r'^(?:saep|[tbg])[\s_-]*', '', s, flags=re.IGNORECASE).strip()
    # 2. Remove suffix EMB, EMP (case-insensitive) only at the end
    s = re.sub(r'[\s_-]*(?:emb|emp)$', '', s, flags=re.IGNORECASE).strip()
    return s

def extract_turma_code(turma_desc):
    if not turma_desc: return ''
    s = str(turma_desc).strip()
    if not s: return ''
    raw_code = s
    for sep in ['-', '–', '—']:
        if sep in raw_code:
            raw_code = raw_code.split(sep)[0].strip()
            break
    return normalize_cod_turma(raw_code)

def is_senai_ited(escola):
    if not escola: return False
    norm = normalize_string(escola)
    return 'senaiited' in norm or norm == 'ited'

def get_unidade_operacional(escola, turma_desc):
    raw_escola = (escola or '').strip()
    if is_senai_ited(raw_escola):
        desc = (turma_desc or '').strip()
        for sep in ['-', '–', '—']:
            if sep in desc:
                parts = desc.split(sep)
                extracted = sep.join(parts[1:]).strip()
                if extracted:
                    return extracted
        return desc or raw_escola
    return raw_escola

# Load real files:
# 1. Aptos (Aba 'Estudante')
wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
ws_aptos = wb_aptos['Estudante']
headers_aptos = [ws_aptos.cell(1, c).value for c in range(1, ws_aptos.max_column+1)]
c_escola = headers_aptos.index('ESCOLA') + 1
c_curso = headers_aptos.index('CURSO') + 1
c_turma = headers_aptos.index('TURMADESCRICAO') + 1

aptos_rows = []
for r in range(2, ws_aptos.max_row+1):
    aptos_rows.append({
        'ESCOLA': ws_aptos.cell(r, c_escola).value,
        'CURSO': ws_aptos.cell(r, c_curso).value,
        'TURMADESCRICAO': ws_aptos.cell(r, c_turma).value,
    })

# 2. CHP (Aba 'Dados', header line 2)
wb_chp = xlrd.open_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
ws_chp = wb_chp.sheet_by_name('Dados')
headers_chp = ws_chp.row_values(1)
c_codturma_chp = headers_chp.index('CODTURMA')
chp_rows = []
for r in range(2, ws_chp.nrows):
    chp_rows.append({
        'CODTURMA': ws_chp.cell_value(r, c_codturma_chp)
    })

# 3. IFTP (Aba 'Dados', header line 1)
wb_iftp = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx', data_only=True)
ws_iftp = wb_iftp['Dados']
headers_iftp = [ws_iftp.cell(1, c).value for c in range(1, ws_iftp.max_column+1)]
c_codturma_iftp = headers_iftp.index('CODTURMA') + 1
iftp_rows = []
for r in range(2, ws_iftp.max_row+1):
    iftp_rows.append({
        'CODTURMA': ws_iftp.cell(r, c_codturma_iftp).value
    })

# 4. Build Indexes
index_chp = defaultdict(int)
for r in chp_rows:
    val = r.get('CODTURMA')
    if val:
        norm = normalize_cod_turma(val)
        if norm:
            index_chp[norm] += 1

index_iftp = defaultdict(int)
for r in iftp_rows:
    val = r.get('CODTURMA')
    if val:
        norm = normalize_cod_turma(val)
        if norm:
            index_iftp[norm] += 1

print(f"Total turmas no índice CHP: {len(index_chp)}")
print(f"Total turmas no índice IFTP: {len(index_iftp)}")

# Group Aptos
class_groups = {}
for r in aptos_rows:
    e = str(r['ESCOLA'] or '').strip()
    c = str(r['CURSO'] or '').strip()
    t = str(r['TURMADESCRICAO'] or '').strip()
    k = (e, c, t)
    if k not in class_groups:
        class_groups[k] = []
    class_groups[k].append(r)

print(f"Total grupos de turma em Aptos: {len(class_groups)}")

# Consolidate
results = []
for (escola, curso, turma_desc), recs in class_groups.items():
    turma_code = extract_turma_code(turma_desc)
    aptos_count = len(recs)
    
    count_chp = index_chp.get(turma_code, 0)
    count_iftp = index_iftp.get(turma_code, 0)
    
    cats = []
    if count_chp > 0:
        cats.append(('CHP', count_chp))
    if count_iftp > 0:
        cats.append(('IFTP', count_iftp))
        
    if not cats:
        cats.append(('CHP', 0))
        
    for cat, alunos in cats:
        inaptos = max(0, alunos - aptos_count)
        results.append({
            'DR': 'SENAI-BA',
            'Unidade': get_unidade_operacional(escola, turma_desc),
            'Curso': curso,
            'Categoria': cat,
            'Modalidade': 'Semipresencial' if is_senai_ited(escola) else 'Presencial',
            'Turma': turma_code,
            'Alunos': alunos,
            'Aptos': aptos_count,
            'Inaptos': inaptos
        })

print(f"\nTotal linhas consolidadas no Macrodado: {len(results)}")
total_alunos = sum(r['Alunos'] for r in results)
total_aptos = sum(r['Aptos'] for r in results)
total_inaptos = sum(r['Inaptos'] for r in results)
print(f"Total Alunos: {total_alunos} | Total Aptos: {total_aptos} | Total Inaptos: {total_inaptos}")

# Inspect some specific IFTP and CHP classes in results
print("\n--- AMOSTRA DE TURMAS CONSOLIDADAS ---")
for r in results[:20]:
    print(f"Turma: {r['Turma']:<18} | Cat: {r['Categoria']:<5} | Alunos: {r['Alunos']:<4} | Aptos: {r['Aptos']:<4} | Inaptos: {r['Inaptos']:<4} | Unidade: {r['Unidade']}")
