import os
import openpyxl

downloads_dir = r"C:\Users\sam.sampaio\Downloads"

def inspect_iftp():
    filename = "Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx"
    path = os.path.join(downloads_dir, filename)
    wb = openpyxl.load_workbook(path, data_only=True)
    print("IFTP Sheets:", wb.sheetnames)
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n==========================================")
        print(f"Sheet '{sname}' (rows={ws.max_row}, cols={ws.max_column})")
        print(f"==========================================")
        for r in range(1, min(15, ws.max_row + 1)):
            row_vals = [ws.cell(r, c).value for c in range(1, min(30, ws.max_column + 1))]
            print(f"Row {r:2d}: {row_vals}")

if __name__ == '__main__':
    inspect_iftp()
