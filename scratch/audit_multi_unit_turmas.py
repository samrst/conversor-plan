import openpyxl
import xlrd
import unicodedata
import re

def normalize_string(s):
    if not s: return ''
    nfd = unicodedata.normalize('NFD', str(s))
    no_accents = ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]', '', no_accents.lower())

def normalize_cod_turma(cod):
    if not cod: return ''
    s = str(cod).strip()
    if not s: return ''
    s = re.sub(r'^(?:saep|[tbg])[\s_-]*', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'[\s_-]*(?:emb|emp)$', '', s, flags=re.IGNORECASE).strip()
    return s

def extract_turma_code(turma_desc):
    if not turma_desc: return ''
    s = str(turma_desc).strip()
    if not s: return ''
    raw_code = s
    for sep in ['-', '–', '—']:
        if sep in raw_code:
            raw_code = raw_code.split(sep)[0].strip()
            break
    return normalize_cod_turma(raw_code)

def is_senai_ited(escola):
    if not escola: return False
    norm = normalize_string(escola)
    return 'senaiited' in norm or norm == 'ited'

def get_unidade_operacional(escola, turma_desc):
    raw_escola = (escola or '').strip()
    if is_senai_ited(raw_escola):
        desc = (turma_desc or '').strip()
        for sep in ['-', '–', '—']:
            if sep in desc:
                parts = desc.split(sep)
                extracted = sep.join(parts[1:]).strip()
                if extracted:
                    return extracted
        return desc or raw_escola
    return raw_escola

# Load real files
wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
ws_aptos = wb_aptos['Estudante']
aptos_unidades = set()
aptos_cursos = set()
aptos_turmas = set()
aptos_combos = set()
turma_units_map = {}

for r in range(2, ws_aptos.max_row + 1):
    esc = ws_aptos.cell(r, 5).value # ESCOLA
    turma_desc = ws_aptos.cell(r, 6).value # TURMADESCRICAO
    curso = ws_aptos.cell(r, 19).value # CURSO
    
    unid = get_unidade_operacional(esc, turma_desc)
    turma_code = extract_turma_code(turma_desc)
    
    aptos_unidades.add((esc, unid, normalize_string(unid)))
    aptos_cursos.add((curso, normalize_string(curso)))
    aptos_turmas.add(turma_code)
    
    u_norm = normalize_string(unid)
    c_norm = normalize_string(curso)
    t_norm = turma_code
    
    if turma_code not in turma_units_map:
        turma_units_map[turma_code] = set()
    turma_units_map[turma_code].add((unid, curso))
    
    aptos_combos.add((u_norm, c_norm, t_norm))

print(f"Total turmas únicas em Aptos: {len(aptos_turmas)}")
print(f"Total combinações (Unidade + Curso + Turma) em Aptos: {len(aptos_combos)}")

# Check for turmas with multiple units in Aptos
multi_unit_turmas = {t: u for t, u in turma_units_map.items() if len(u) > 1}
print(f"\nTurmas com MAIS DE UMA unidade/curso em Aptos: {len(multi_unit_turmas)}")
for t, units in list(multi_unit_turmas.items())[:10]:
    print(f"  Turma {t}:")
    for u, c in units:
        print(f"    - Unidade: {u} | Curso: {c}")

# Load CHP
wb_chp = xlrd.open_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
ws_chp = wb_chp.sheet_by_name('Dados')
chp_combos = set()
chp_turma_units = {}
for r in range(2, ws_chp.nrows):
    unid = ws_chp.cell_value(r, 0)
    curso = ws_chp.cell_value(r, 2)
    cod = ws_chp.cell_value(r, 4)
    if cod:
        t_norm = normalize_cod_turma(cod)
        u_norm = normalize_string(unid)
        c_norm = normalize_string(curso)
        chp_combos.add((u_norm, c_norm, t_norm))
        if t_norm not in chp_turma_units:
            chp_turma_units[t_norm] = set()
        chp_turma_units[t_norm].add((unid, curso))

print(f"\nTotal combinações (Unidade + Curso + Turma) em CHP: {len(chp_combos)}")
multi_unit_chp = {t: u for t, u in chp_turma_units.items() if len(u) > 1}
print(f"Turmas com MAIS DE UMA unidade/curso em CHP: {len(multi_unit_chp)}")
for t, units in list(multi_unit_chp.items())[:10]:
    print(f"  Turma {t}:")
    for u, c in units:
        print(f"    - Unidade: {u} | Curso: {c}")
