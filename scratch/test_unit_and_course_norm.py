import openpyxl
import xlrd
import unicodedata
import re

def normalize_string(s):
    if not s: return ''
    nfd = unicodedata.normalize('NFD', str(s))
    no_accents = ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]', '', no_accents.lower())

def normalize_unidade(unidade):
    if not unidade: return ''
    s = normalize_string(unidade)
    
    # Remove generic school prefixes
    s = re.sub(r'^(?:senai|cfp|escola|unidade|centro|cetind)+', '', s)
    
    # Canonical mappings for common Bahia unit variations
    canonical = {
        'feiradesantana': 'feira',
        'feiradesantanabest': 'feira',
        'feiraserrinha': 'feira',
        'vitoriadaconquista': 'conquista',
        'luiseduardomagalhaes': 'lem',
        'luiseduardo': 'lem',
        'salvadorcimatec': 'cimatec',
        'salvadordendezeiros': 'dendezeiros',
    }
    for k, v in canonical.items():
        if s.startswith(k) or s == k:
            return v
    return s

def normalize_curso(curso):
    if not curso: return ''
    s = normalize_string(curso)
    # Remove prefix "tecnicoem", "tecnicode", "curso", etc. if desired or keep clean
    s = re.sub(r'^(?:tecnicoem|tecnicode|cursobtecnico|curso)+', '', s)
    return s

print("--- TESTE DE NORMALIZAÇÃO DE UNIDADES ---")
test_units_aptos = [
    'ALAGOINHAS', 'CAMAÇARI', 'CFP Candeias', 'CFP Teixeira de Freitas',
    'CONQUISTA', 'DENDEZEIROS', 'FEIRA', 'FEIRA(SERRINHA)', 'ILHÉUS',
    'JUAZEIRO', 'LAURO DE FREITAS', 'LEM', 'SENAI Alagoinhas',
    'SENAI CIMATEC', 'SENAI Camaçari', 'SENAI Dendezeiros',
    'SENAI Feira de Santana', 'SENAI Juazeiro', 'SENAI Lauro de Freitas',
    'SENAI Vitória da Conquista', 'TEIXEIRA DE FREITAS'
]

test_units_matriculados = [
    'ALAGOINHAS', 'CAMAÇARI', 'CIMATEC', 'CONQUISTA', 'DENDEZEIROS',
    'FEIRA', 'ILHÉUS', 'JUAZEIRO', 'LAURO DE FREITAS', 'LEM',
    'TEIXEIRA DE FREITAS', 'CANDEIAS'
]

print("Aptos Unidades normalizadas:")
for u in test_units_aptos:
    print(f"  {u:<30} -> {normalize_unidade(u)}")

print("\nMatriculados Unidades normalizadas:")
for u in test_units_matriculados:
    print(f"  {u:<30} -> {normalize_unidade(u)}")

print("\n--- TESTE DE CURSOS EM MATRICULADOS E APTOS ---")
wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
ws_aptos = wb_aptos['Estudante']
aptos_cursos = set(ws_aptos.cell(r, 19).value for r in range(2, ws_aptos.max_row + 1) if ws_aptos.cell(r, 19).value)

wb_chp = xlrd.open_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
ws_chp = wb_chp.sheet_by_name('Dados')
chp_cursos = set(ws_chp.cell_value(r, 2) for r in range(2, ws_chp.nrows) if ws_chp.cell_value(r, 2))

wb_iftp = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx', data_only=True)
ws_iftp = wb_iftp['Dados']
iftp_cursos = set(ws_iftp.cell(r, 8).value for r in range(2, ws_iftp.max_row + 1) if ws_iftp.cell(r, 8).value)

print("Aptos Cursos:")
for c in sorted(list(aptos_cursos)):
    print(f"  {c:<45} -> {normalize_curso(c)}")

print("\nCHP Cursos:")
for c in sorted(list(chp_cursos)):
    print(f"  {c:<45} -> {normalize_curso(c)}")

print("\nIFTP Cursos:")
for c in sorted(list(iftp_cursos)):
    print(f"  {c:<45} -> {normalize_curso(c)}")
