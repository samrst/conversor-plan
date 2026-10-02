import openpyxl
import xlrd
import os
import re
import unicodedata
from collections import defaultdict

def normalize_string(s):
    if not s: return ''
    nfd = unicodedata.normalize('NFD', str(s))
    no_accents = ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]', '', no_accents.lower())

def normalize_cod_turma(cod):
    if not cod: return ''
    s = str(cod).strip()
    if not s: return ''
    s = re.sub(r'^(?:saep|[tbg])[\s_-]*', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'[\s_-]*(?:emb|emp)$', '', s, flags=re.IGNORECASE).strip()
    return s

def extract_turma_code(turma_desc):
    if not turma_desc: return ''
    s = str(turma_desc).strip()
    if not s: return ''
    raw_code = s
    for sep in ['-', '–', '—']:
        if sep in raw_code:
            raw_code = raw_code.split(sep)[0].strip()
            break
    return normalize_cod_turma(raw_code)

def is_senai_ited(escola):
    if not escola: return False
    norm = normalize_string(escola)
    return 'senaiited' in norm or norm == 'ited'

def get_unidade_operacional(escola, turma_desc):
    raw_escola = (escola or '').strip()
    if is_senai_ited(raw_escola):
        desc = (turma_desc or '').strip()
        for sep in ['-', '–', '—']:
            if sep in desc:
                parts = desc.split(sep)
                extracted = sep.join(parts[1:]).strip()
                if extracted:
                    return extracted
        return desc or raw_escola
    return raw_escola

def normalize_unidade(unidade):
    if not unidade: return ''
    s = normalize_string(unidade)
    s = re.sub(r'^(?:senai|cfp|escola|unidade|centro|cetind)+', '', s)
    canonical = {
        'feiradesantana': 'feira',
        'feiradesantanabest': 'feira',
        'feiraserrinha': 'feira',
        'vitoriadaconquista': 'conquista',
        'luiseduardomagalhaes': 'lem',
        'luiseduardo': 'lem',
        'salvadorcimatec': 'cimatec',
        'salvadordendezeiros': 'dendezeiros'
    }
    for k, v in canonical.items():
        if s.startswith(k) or s == k:
            return v
    return s

def normalize_curso(curso):
    if not curso: return ''
    s = normalize_string(curso)
    s = re.sub(r'^(?:tecnicoem|tecnicode|cursobtecnico|curso)+', '', s)
    return s

def make_composite_key(unidade, curso, turma_code):
    u = normalize_unidade(unidade)
    c = normalize_curso(curso)
    t = normalize_cod_turma(turma_code)
    return f"{u}___{c}___{t}"

# 1. Load Aptos
wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
ws_aptos = wb_aptos['Estudante']
headers_aptos = [ws_aptos.cell(1, c).value for c in range(1, ws_aptos.max_column + 1)]
rows_aptos = []
for r in range(2, ws_aptos.max_row + 1):
    vals = [ws_aptos.cell(r, c).value for c in range(1, ws_aptos.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_aptos.append(dict(zip(headers_aptos, vals)))

# 2. Load CHP
wb_chp = xlrd.open_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
ws_chp = wb_chp.sheet_by_name('Dados')
headers_chp = [ws_chp.cell_value(1, c) for c in range(ws_chp.ncols)]
rows_chp = []
for r in range(2, ws_chp.nrows):
    vals = [ws_chp.cell_value(r, c) for c in range(ws_chp.ncols)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_chp.append(dict(zip(headers_chp, vals)))

# 3. Load IFTP
wb_iftp = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx', data_only=True)
ws_iftp = wb_iftp['Dados']
headers_iftp = [ws_iftp.cell(1, c).value for c in range(1, ws_iftp.max_column + 1)]
rows_iftp = []
for r in range(2, ws_iftp.max_row + 1):
    vals = [ws_iftp.cell(r, c).value for c in range(1, ws_iftp.max_column + 1)]
    if any(v is not None and str(v).strip() != '' for v in vals):
        rows_iftp.append(dict(zip(headers_iftp, vals)))

print(f"=== TOTAL RECORDS IN SPREADSHEETS ===")
print(f"Aptos (Estudantes): {len(rows_aptos)}")
print(f"CHP (Estudantes Matriculados): {len(rows_chp)}")
print(f"IFTP (Estudantes Matriculados): {len(rows_iftp)}")
print(f"Grand Total Matriculados in Files: {len(rows_chp) + len(rows_iftp)}")

# Inspect 96225, 96226, 93629 across files
print(f"\n=== INSPECTING UNMATCHED CODES ===")
for target_code in ['96225', '96226', '93629']:
    print(f"\n--- Code: {target_code} ---")
    
    # In Aptos:
    aptos_hits = [r for r in rows_aptos if extract_turma_code(r.get('TURMADESCRICAO')) == target_code]
    print(f"Aptos rows: {len(aptos_hits)}")
    for r in aptos_hits:
        print(f"  Aptos: Escola='{r.get('ESCOLA')}', Curso='{r.get('CURSO')}', TurmaDesc='{r.get('TURMADESCRICAO')}'")
        
    # In CHP:
    chp_hits = [r for r in rows_chp if normalize_cod_turma(r.get('CODTURMA')) == target_code]
    print(f"CHP rows: {len(chp_hits)}")
    for r in chp_hits[:5]:
        print(f"  CHP: Unidade='{r.get('UNIDADE') or r.get('ESCOLA')}', Curso='{r.get('CURSO')}', CodTurma='{r.get('CODTURMA')}'")
    if len(chp_hits) > 5:
        print(f"  ... and {len(chp_hits)-5} more in CHP")
        
    # In IFTP:
    iftp_hits = [r for r in rows_iftp if normalize_cod_turma(r.get('CODTURMA')) == target_code]
    print(f"IFTP rows: {len(iftp_hits)}")
    for r in iftp_hits[:5]:
        print(f"  IFTP: Unidade='{r.get('UNIDADE') or r.get('ESCOLA')}', Curso='{r.get('CURSO')}', CodTurma='{r.get('CODTURMA')}'")

# Check all classes in CHP that were NOT associated
class_groups = defaultdict(list)
for r in rows_aptos:
    raw_escola = str(r.get('ESCOLA', '')).strip()
    raw_curso = str(r.get('CURSO', '')).strip()
    raw_turma = str(r.get('TURMADESCRICAO', '')).strip()
    unidade_op = get_unidade_operacional(raw_escola, raw_turma)
    class_groups[(unidade_op, raw_curso, raw_turma, raw_escola)].append(r)

aptos_comp_keys = set()
for (unidade_op, curso, turma_desc, escola_raw), recs in class_groups.items():
    turma_code = extract_turma_code(turma_desc)
    comp_key = make_composite_key(unidade_op, curso, turma_code)
    aptos_comp_keys.add(comp_key)

chp_matched_students = 0
chp_unmatched_students = 0
chp_unmatched_by_class = defaultdict(list)

for r in rows_chp:
    c = r.get('CODTURMA')
    u = r.get('UNIDADE') or r.get('ESCOLA') or ''
    cur = r.get('CURSO') or ''
    comp = make_composite_key(u, cur, c)
    if comp in aptos_comp_keys:
        chp_matched_students += 1
    else:
        chp_unmatched_students += 1
        chp_unmatched_by_class[(u, cur, c)].append(r)

iftp_matched_students = 0
iftp_unmatched_students = 0
iftp_unmatched_by_class = defaultdict(list)

for r in rows_iftp:
    c = r.get('CODTURMA')
    u = r.get('UNIDADE') or r.get('ESCOLA') or ''
    cur = r.get('CURSO') or ''
    comp = make_composite_key(u, cur, c)
    if comp in aptos_comp_keys:
        iftp_matched_students += 1
    else:
        iftp_unmatched_students += 1
        iftp_unmatched_by_class[(u, cur, c)].append(r)

print(f"\n=== DETAILED RECONCILIATION OF STUDENTS ===")
print(f"CHP:")
print(f"  Total no arquivo: {len(rows_chp)}")
print(f"  Associados (em turmas avaliadas no Aptos): {chp_matched_students}")
print(f"  Não associados: {chp_unmatched_students}")

print(f"\nIFTP:")
print(f"  Total no arquivo: {len(rows_iftp)}")
print(f"  Associados (em turmas avaliadas no Aptos): {iftp_matched_students}")
print(f"  Não associados: {iftp_unmatched_students}")

print(f"\nTOTAL GERAL:")
print(f"  Total nos arquivos: {len(rows_chp) + len(rows_iftp)}")
print(f"  Total associado: {chp_matched_students + iftp_matched_students}")
print(f"  Total não associado: {chp_unmatched_students + iftp_unmatched_students}")
print(f"  Soma fecha? {(chp_matched_students + iftp_matched_students) + (chp_unmatched_students + iftp_unmatched_students) == len(rows_chp) + len(rows_iftp)}")

print(f"\n=== UNMATCHED CLASSES IN CHP (CLASSES IN CHP NOT IN APTOS) ===")
for (u, cur, c), recs in chp_unmatched_by_class.items():
    print(f"  Unidade='{u}', Curso='{cur}', Turma='{c}' -> {len(recs)} alunos")

print(f"\n=== UNMATCHED CLASSES IN IFTP (CLASSES IN IFTP NOT IN APTOS) ===")
for (u, cur, c), recs in iftp_unmatched_by_class.items():
    print(f"  Unidade='{u}', Curso='{cur}', Turma='{c}' -> {len(recs)} alunos")
