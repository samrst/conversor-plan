import openpyxl
import xlrd
import unicodedata
import re

def normalize_string(s):
    if not s: return ''
    nfd = unicodedata.normalize('NFD', str(s))
    no_accents = ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]', '', no_accents.lower())

def is_senai_ited(escola):
    if not escola: return False
    norm = normalize_string(escola)
    return 'senaiited' in norm or norm == 'ited'

def get_unidade_operacional(escola, turma_desc):
    raw_escola = (escola or '').strip()
    if is_senai_ited(raw_escola):
        desc = (turma_desc or '').strip()
        for sep in ['-', '–', '—']:
            if sep in desc:
                parts = desc.split(sep)
                extracted = sep.join(parts[1:]).strip()
                if extracted:
                    return extracted
        return desc or raw_escola
    return raw_escola

# Load all units from Aptos
wb_aptos = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\SAEP 2026-2_28_08_2026_BA.xlsx', data_only=True)
ws_aptos = wb_aptos['Estudante']
aptos_units = set()
for r in range(2, ws_aptos.max_row + 1):
    esc = ws_aptos.cell(r, 5).value
    desc = ws_aptos.cell(r, 6).value
    u_op = get_unidade_operacional(esc, desc)
    aptos_units.add(u_op)

# Load all units from CHP
wb_chp = xlrd.open_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls')
ws_chp = wb_chp.sheet_by_name('Dados')
chp_units = set()
for r in range(2, ws_chp.nrows):
    u = ws_chp.cell_value(r, 0)
    if u:
        chp_units.add(u)

# Load all units from IFTP
wb_iftp = openpyxl.load_workbook(r'C:\Users\sam.sampaio\Downloads\Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx', data_only=True)
ws_iftp = wb_iftp['Dados']
iftp_units = set()
for r in range(2, ws_iftp.max_row + 1):
    u = ws_iftp.cell(r, 1).value
    if u:
        iftp_units.add(u)

print("=== UNIDADES EM APTOS ===")
for u in sorted(list(aptos_units)):
    print("  Aptos:", repr(u), "-> norm:", normalize_string(u))

print("\n=== UNIDADES EM CHP ===")
for u in sorted(list(chp_units)):
    print("  CHP:", repr(u), "-> norm:", normalize_string(u))

print("\n=== UNIDADES EM IFTP ===")
for u in sorted(list(iftp_units)):
    print("  IFTP:", repr(u), "-> norm:", normalize_string(u))
