import os
import openpyxl
import xlrd
import pandas as pd

def inspect_file(filepath):
    print(f"\n=======================================================")
    print(f"FILE: {os.path.basename(filepath)}")
    print(f"=======================================================")
    if filepath.endswith('.xlsx'):
        wb = openpyxl.load_workbook(filepath, data_only=True)
        print("Sheets:", wb.sheetnames)
        for sname in wb.sheetnames[:3]:
            ws = wb[sname]
            print(f"\n--- Sheet: {sname} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
            # Headers (row 1 or find header row)
            for r in range(1, min(6, ws.max_row + 1)):
                row_vals = [ws.cell(r, c).value for c in range(1, min(20, ws.max_column + 1))]
                print(f"Row {r}: {row_vals}")
    elif filepath.endswith('.xls'):
        wb = xlrd.open_workbook(filepath)
        print("Sheets:", wb.sheet_names())
        for sname in wb.sheet_names()[:3]:
            ws = wb.sheet_by_name(sname)
            print(f"\n--- Sheet: {sname} (nrows={ws.nrows}, ncols={ws.ncols}) ---")
            for r in range(min(5, ws.nrows)):
                print(f"Row {r+1}: {ws.row_values(r)[:20]}")

if __name__ == '__main__':
    downloads_dir = r"C:\Users\sam.sampaio\Downloads"
    files = [
        "Alunos matriculados 2026 IFTP - SAEP 2026.2.xlsx",
        "Alunos Matriculados 26.1 CHP - SAEP 2026.2.xls",
        "SAEP 2026-2_28_08_2026_BA.xlsx",
        "SAEP 2026-2_01_09_2026_BA - FINAL.xlsx",
        "RelatorioParticipacaoTurmasProvas_2026-06-22_14-42-12.xlsx"
    ]
    for f in files:
        path = os.path.join(downloads_dir, f)
        if os.path.exists(path):
            inspect_file(path)
        else:
            print(f"File not found: {path}")
